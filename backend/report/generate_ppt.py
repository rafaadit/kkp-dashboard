#!/usr/bin/env python3
"""
PHASE 7 — Generate laporan PPT dari data raw_exim (python-pptx).

Struktur mengikuti deck referensi "Draft_EKSIM" (11 slide):
  0  Cover
  1  Kinerja Ekspor Jan-Jun 2022-2026 (line chart)
  2  Ekspor-Impor-Neraca (bar chart)
  3  Negara Tujuan Ekspor (tabel)
  4  Komoditas Utama Ekspor (tabel)
  5  Pasar Utama Ekspor (pelabuhan) (tabel)
  6  Produk Utama Ekspor (tabel komoditas_5_2026)
  7  Kinerja Impor Jan-Jun 2022-2026 (line chart)
  8  Negara Asal Impor (tabel + total)
  9  Komoditas Utama Impor (tabel)
 10  Catatan/penutup

Nilai dari DB (raw_exim). Output ke PATH_REPORT (default data/reports/).
Tercatat di automation_jobs (job_type='report_ppt').

Usage:
  .venv/bin/python backend/report/generate_ppt.py --tahun 2026 --bulan-awal 1 --bulan-akhir 6
  .venv/bin/python backend/report/generate_ppt.py --tahun 2026 --bulan-awal 1 --bulan-akhir 6 --output myreport.pptx
"""
import argparse
import json
import os
import sys
import uuid
from datetime import datetime

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
if ROOT not in sys.path:
    sys.path.insert(0, ROOT)

from pptx import Presentation  # noqa: E402
from pptx.chart.data import CategoryChartData  # noqa: E402
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION  # noqa: E402
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR  # noqa: E402
from pptx.util import Emu, Inches, Pt  # noqa: E402

from backend.db import get_connection  # noqa: E402

NAVY = RGBColor(0x14, 0x37, 0x5E)
GOLD = RGBColor(0xF2, 0xA9, 0x00)
BLUE = RGBColor(0x0B, 0x6B, 0xCB)
GREY = RGBColor(0x6B, 0x76, 0x84)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT = RGBColor(0xEE, 0xF3, 0xF8)
RED = RGBColor(0xB9, 0x1C, 0x1C)
GREEN = RGBColor(0x15, 0x80, 0x3D)

TITLE = "KEMENTERIAN KELAUTAN DAN PERIKANAN"
REPORT_DIR = os.environ.get("PATH_REPORT", os.path.join(ROOT, "data", "reports"))

BULAN_ID = ["", "Januari", "Februari", "Maret", "April", "Mei", "Juni",
            "Juli", "Agustus", "September", "Oktober", "November", "Desember"]


def fmt_usd(v, satuan="juta"):
    if v is None:
        return "-"
    v = float(v)
    if satuan == "miliar":
        return f"{v/1e9:,.2f}"
    if satuan == "juta":
        return f"{v/1e6:,.2f}"
    return f"{v:,.0f}"


def fmt_num(v):
    return f"{float(v):,.2f}" if v is not None else "-"


def q(conn, sql, params=()):
    cur = conn.cursor()
    cur.execute(sql, params)
    return cur.fetchall()


# --------------------------------------------------------------------------
# Data
# --------------------------------------------------------------------------
def collect_data(conn, tahun, b1, b2):
    pa = "%d%02d" % (tahun, b1)
    pb = "%d%02d" % (tahun, b2)

    per_tahun = q(conn, """
        SELECT tahun,
               COALESCE(SUM(CASE WHEN exim_type='ekspor' THEN nil_usd END),0) AS ekspor,
               COALESCE(SUM(CASE WHEN exim_type='impor'  THEN nil_usd END),0) AS impor
        FROM raw_exim
        WHERE bulan BETWEEN %s AND %s AND tahun BETWEEN 2022 AND %s
        GROUP BY tahun ORDER BY tahun
    """, (b1, b2, tahun))

    total = q(conn, """
        SELECT
          COALESCE(SUM(CASE WHEN exim_type='ekspor' THEN nil_usd END),0) AS ekspor_usd,
          COALESCE(SUM(CASE WHEN exim_type='ekspor' THEN vol_kg END),0) AS ekspor_kg,
          COALESCE(SUM(CASE WHEN exim_type='impor'  THEN nil_usd END),0) AS impor_usd,
          COALESCE(SUM(CASE WHEN exim_type='impor'  THEN vol_kg END),0) AS impor_kg
        FROM raw_exim WHERE tahun_bulan BETWEEN %s AND %s
    """, (pa, pb))[0]

    negara = q(conn, """
        SELECT negara, kode_negara, COALESCE(SUM(nil_usd),0) nilai, COALESCE(SUM(vol_kg),0) vol
        FROM raw_exim WHERE exim_type='ekspor' AND tahun_bulan BETWEEN %s AND %s
        GROUP BY negara, kode_negara ORDER BY nilai DESC LIMIT 10
    """, (pa, pb))

    komoditas = q(conn, """
        SELECT komoditas_5_2026 komoditas, COALESCE(SUM(nil_usd),0) nilai, COALESCE(SUM(vol_kg),0) vol
        FROM raw_exim WHERE exim_type='ekspor' AND tahun_bulan BETWEEN %s AND %s
              AND NULLIF(komoditas_5_2026,'') IS NOT NULL
        GROUP BY komoditas_5_2026 ORDER BY nilai DESC LIMIT 10
    """, (pa, pb))

    pasar = q(conn, """
        SELECT pelabuhan_muat_bongkar pelabuhan, COALESCE(SUM(nil_usd),0) nilai
        FROM raw_exim WHERE exim_type='ekspor' AND tahun_bulan BETWEEN %s AND %s
              AND NULLIF(pelabuhan_muat_bongkar,'') IS NOT NULL
        GROUP BY pelabuhan_muat_bongkar ORDER BY nilai DESC LIMIT 10
    """, (pa, pb))

    negara_impor = q(conn, """
        SELECT negara, kode_negara, COALESCE(SUM(nil_usd),0) nilai
        FROM raw_exim WHERE exim_type='impor' AND tahun_bulan BETWEEN %s AND %s
        GROUP BY negara, kode_negara ORDER BY nilai DESC LIMIT 10
    """, (pa, pb))

    komoditas_impor = q(conn, """
        SELECT komoditas_5_2026 komoditas, COALESCE(SUM(nil_usd),0) nilai, COALESCE(SUM(vol_kg),0) vol
        FROM raw_exim WHERE exim_type='impor' AND tahun_bulan BETWEEN %s AND %s
              AND NULLIF(komoditas_5_2026,'') IS NOT NULL
        GROUP BY komoditas_5_2026 ORDER BY nilai DESC LIMIT 10
    """, (pa, pb))

    prev = q(conn, """
        SELECT COALESCE(SUM(nil_usd),0) AS v FROM raw_exim
        WHERE exim_type='impor' AND tahun=%s AND bulan BETWEEN %s AND %s
    """, (tahun - 1, b1, b2))[0]["v"]

    return {
        "tahun": tahun, "b1": b1, "b2": b2,
        "per_tahun": per_tahun, "total": total, "negara": negara,
        "komoditas": komoditas, "pasar": pasar,
        "negara_impor": negara_impor, "komoditas_impor": komoditas_impor,
        "impor_prev": prev,
    }


# --------------------------------------------------------------------------
# Helpers PPT
# --------------------------------------------------------------------------
def _blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def _rect(slide, x, y, w, h, color, line=False):
    from pptx.enum.shapes import MSO_SHAPE
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = color
    if not line:
        sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def _text(slide, x, y, w, h, text, size=18, color=NAVY, bold=False, align=PP_ALIGN.LEFT, font="Calibri"):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    r.font.name = font
    return tb


def _header(slide, prs, nomor, subtitle):
    _rect(slide, 0, 0, prs.slide_width, Inches(0.75), NAVY)
    _text(slide, Inches(0.3), Inches(0.08), Inches(10), Inches(0.35),
          TITLE, size=13, color=WHITE, bold=True)
    _text(slide, Inches(0.3), Inches(0.36), Inches(11), Inches(0.32),
          subtitle, size=11, color=RGBColor(0xCF, 0xE4, 0xF7))
    _text(slide, Emu(prs.slide_width) - Inches(1.2), Inches(0.14), Inches(1), Inches(0.5),
          str(nomor), size=20, color=GOLD, bold=True, align=PP_ALIGN.RIGHT)


def _table(slide, headers, rows, x=Inches(0.4), y=Inches(1.1), w=Inches(12.4), h=Inches(5.3)):
    n_rows = len(rows) + 1
    n_cols = len(headers)
    gf = slide.shapes.add_table(n_rows, n_cols, x, y, w, h)
    table = gf.table
    for j, htxt in enumerate(headers):
        cell = table.cell(0, j)
        cell.text = htxt
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY
        for p in cell.text_frame.paragraphs:
            p.alignment = PP_ALIGN.CENTER
            for r in p.runs:
                r.font.size = Pt(11)
                r.font.bold = True
                r.font.color.rgb = WHITE
    for i, row in enumerate(rows, start=1):
        for j, val in enumerate(row):
            cell = table.cell(i, j)
            cell.text = str(val)
            cell.fill.solid()
            cell.fill.fore_color.rgb = WHITE if i % 2 else LIGHT
            for p in cell.text_frame.paragraphs:
                p.alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.RIGHT
                for r in p.runs:
                    r.font.size = Pt(10)
                    r.font.color.rgb = NAVY
    return table


def _chart(slide, chart_type, categories, series, x, y, w, h, number_format='#,##0'):
    cd = CategoryChartData()
    cd.categories = categories
    for name, values in series:
        cd.add_series(name, values)
    gf = slide.shapes.add_chart(chart_type, x, y, w, h, cd)
    chart = gf.chart
    chart.has_legend = len(series) > 1
    if chart.has_legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
    try:
        chart.value_axis.tick_labels.number_format = number_format
    except Exception:
        pass
    return chart


# --------------------------------------------------------------------------
# Slides
# --------------------------------------------------------------------------
def slide_cover(prs, d):
    s = _blank(prs)
    _rect(s, 0, 0, prs.slide_width, prs.slide_height, NAVY)
    _rect(s, 0, Emu(int(prs.slide_height)) - Inches(0.3), prs.slide_width, Inches(0.3), GOLD)
    _text(s, Inches(0.8), Inches(1.2), Inches(11), Inches(0.6), TITLE, size=22, color=WHITE, bold=True)
    _text(s, Inches(0.8), Inches(2.2), Inches(11), Inches(0.9), "NERACA", size=48, color=GOLD, bold=True)
    _text(s, Inches(0.8), Inches(3.1), Inches(11), Inches(0.9), "PERDAGANGAN", size=48, color=WHITE, bold=True)
    _text(s, Inches(0.8), Inches(4.0), Inches(11), Inches(0.9), "LUAR NEGERI", size=48, color=WHITE, bold=True)
    periode = f"PERIODE {BULAN_ID[d['b1']].upper()} – {BULAN_ID[d['b2']].upper()} {d['tahun']}"
    _text(s, Inches(0.8), Inches(5.2), Inches(11), Inches(0.5), periode, size=18, color=WHITE)
    _text(s, Inches(0.8), Inches(6.0), Inches(11), Inches(0.4),
          f"Dibuat otomatis: {datetime.now().strftime('%d %b %Y %H:%M')} | Sumber: raw_exim (BPS)",
          size=11, color=RGBColor(0x9F, 0xC0, 0xDD))


def slide_ekspor_kinerja(prs, d):
    s = _blank(prs)
    _header(s, prs, 2, f"KINERJA EKSPOR HASIL PERIKANAN — PERIODE {BULAN_ID[d['b1']].upper()}–{BULAN_ID[d['b2']].upper()} 2022-{d['tahun']}")
    cats = [str(r["tahun"]) for r in d["per_tahun"]]
    vals = [float(r["ekspor"]) / 1e9 for r in d["per_tahun"]]
    _chart(s, XL_CHART_TYPE.LINE_MARKERS, cats, [("Ekspor (USD miliar)", vals)],
           Inches(0.5), Inches(1.1), Inches(12.2), Inches(5.2), '#,##0.00')


def slide_neraca(prs, d):
    s = _blank(prs)
    _header(s, prs, 3, f"EKSPOR-IMPOR-NERACA PRODUK PERIKANAN (USD MILIAR) — {BULAN_ID[d['b1']].upper()}–{BULAN_ID[d['b2']].upper()} 2022-{d['tahun']}")
    cats = [str(r["tahun"]) for r in d["per_tahun"]]
    eks = [float(r["ekspor"]) / 1e9 for r in d["per_tahun"]]
    imp = [float(r["impor"]) / 1e9 for r in d["per_tahun"]]
    ner = [e - i for e, i in zip(eks, imp)]
    _chart(s, XL_CHART_TYPE.COLUMN_CLUSTERED, cats,
           [("Ekspor", eks), ("Impor", imp), ("Neraca", ner)],
           Inches(0.5), Inches(1.1), Inches(12.2), Inches(5.2), '#,##0.00')


def slide_negara_tujuan(prs, d):
    s = _blank(prs)
    _header(s, prs, 4, f"NEGARA TUJUAN EKSPOR — {BULAN_ID[d['b1']].upper()}–{BULAN_ID[d['b2']].upper()} {d['tahun']}")
    rows = [[r["negara"] or r["kode_negara"], fmt_usd(r["nilai"]), fmt_num(r["vol"] / 1000 if r["vol"] else 0)]
            for r in d["negara"]]
    _table(s, ["Negara", "Nilai (USD juta)", "Volume (ton)"], rows)


def slide_komoditas(prs, d):
    s = _blank(prs)
    _header(s, prs, 5, f"KOMODITAS UTAMA EKSPOR — {BULAN_ID[d['b1']].upper()}–{BULAN_ID[d['b2']].upper()} {d['tahun']}")
    rows = [[r["komoditas"], fmt_usd(r["nilai"]), fmt_num(r["vol"] / 1000 if r["vol"] else 0)]
            for r in d["komoditas"]]
    _table(s, ["Komoditas", "Nilai (USD juta)", "Volume (ton)"], rows)


def slide_pasar(prs, d):
    s = _blank(prs)
    _header(s, prs, 6, f"PASAR/ PELABUHAN UTAMA EKSPOR — {BULAN_ID[d['b1']].upper()}–{BULAN_ID[d['b2']].upper()} {d['tahun']}")
    rows = [[r["pelabuhan"], fmt_usd(r["nilai"])] for r in d["pasar"]]
    _table(s, ["Pelabuhan Muat", "Nilai (USD juta)"], rows)


def slide_produk(prs, d):
    s = _blank(prs)
    _header(s, prs, 7, f"PRODUK UTAMA EKSPOR — {BULAN_ID[d['b1']].upper()}–{BULAN_ID[d['b2']].upper()} {d['tahun']}")
    rows = [[r["komoditas"], fmt_usd(r["nilai"]), fmt_num(r["vol"])] for r in d["komoditas"]]
    _table(s, ["Produk (kelompok komoditas)", "Nilai (USD juta)", "Volume (kg)"], rows)


def slide_impor_kinerja(prs, d):
    s = _blank(prs)
    _header(s, prs, 8, f"KINERJA IMPOR HASIL PERIKANAN — PERIODE {BULAN_ID[d['b1']].upper()}–{BULAN_ID[d['b2']].upper()} 2022-{d['tahun']}")
    cats = [str(r["tahun"]) for r in d["per_tahun"]]
    vals = [float(r["impor"]) / 1e9 for r in d["per_tahun"]]
    _chart(s, XL_CHART_TYPE.LINE_MARKERS, cats, [("Impor (USD miliar)", vals)],
           Inches(0.5), Inches(1.1), Inches(12.2), Inches(5.2), '#,##0.00')


def slide_negara_impor(prs, d):
    s = _blank(prs)
    _header(s, prs, 9, f"NEGARA ASAL IMPOR — {BULAN_ID[d['b1']].upper()}–{BULAN_ID[d['b2']].upper()} {d['tahun']}")
    _text(s, Inches(0.4), Inches(0.95), Inches(8), Inches(0.4),
          f"TOTAL IMPOR: USD {fmt_usd(d['total']['impor_usd'])} JUTA", size=14, color=RED, bold=True)
    rows = [[r["negara"] or r["kode_negara"], fmt_usd(r["nilai"])] for r in d["negara_impor"]]
    _table(s, ["Negara Asal", "Nilai (USD juta)"], rows, y=Inches(1.5), h=Inches(4.9))


def slide_komoditas_impor(prs, d):
    s = _blank(prs)
    _header(s, prs, 10, f"KOMODITAS UTAMA IMPOR — {BULAN_ID[d['b1']].upper()}–{BULAN_ID[d['b2']].upper()} {d['tahun']}")
    rows = [[r["komoditas"], fmt_usd(r["nilai"]), fmt_num(r["vol"] / 1000 if r["vol"] else 0)]
            for r in d["komoditas_impor"]]
    _table(s, ["Komoditas", "Nilai (USD juta)", "Volume (ton)"], rows)


def slide_penutup(prs, d):
    s = _blank(prs)
    _rect(s, 0, 0, prs.slide_width, prs.slide_height, NAVY)
    _text(s, Inches(0.8), Inches(2.5), Inches(11), Inches(0.8), "CATATAN", size=30, color=GOLD, bold=True)
    notes = [
        "Angka bersumber dari tabel raw_exim (hasil olahan data BPS).",
        "Nilai ekspor/impor dalam USD; volume sesuai satuan asal (kg).",
        "Baris berstatus needs_validation berarti kurs/nilai belum final.",
    ]
    y = Inches(3.5)
    for n in notes:
        _text(s, Inches(0.8), y, Inches(11), Inches(0.5), "• " + n, size=14, color=WHITE)
        y += Inches(0.55)


# --------------------------------------------------------------------------
# Main
# --------------------------------------------------------------------------
def generate(tahun, b1, b2, output=None, created_by=None):
    conn = get_connection()
    cur = conn.cursor()
    corr = uuid.uuid4().hex
    cur.execute(
        """INSERT INTO automation_jobs (job_type, status, params_json, attempts, correlation_id, started_at, created_by)
           VALUES ('report_ppt', 'RUNNING', %s, 1, %s, NOW(), %s)""",
        (json.dumps({"tahun": tahun, "bulan_awal": b1, "bulan_akhir": b2}), corr, created_by),
    )
    job_id = cur.lastrowid
    conn.commit()

    def log(level, msg, data=None):
        c2 = conn.cursor()
        c2.execute("INSERT INTO automation_job_logs (id_job, level, message, data_json) VALUES (%s,%s,%s,%s)",
                   (job_id, level, msg, json.dumps(data) if data else None))
        conn.commit()

    try:
        data = collect_data(conn, tahun, b1, b2)
        if not data["per_tahun"]:
            raise RuntimeError("tidak ada data untuk periode ini")
        log("info", "data terkumpul", {"tahun": tahun, "b1": b1, "b2": b2})

        prs = Presentation()
        prs.slide_width = Emu(12192000)
        prs.slide_height = Emu(6858000)
        slide_cover(prs, data)
        slide_ekspor_kinerja(prs, data)
        slide_neraca(prs, data)
        slide_negara_tujuan(prs, data)
        slide_komoditas(prs, data)
        slide_pasar(prs, data)
        slide_produk(prs, data)
        slide_impor_kinerja(prs, data)
        slide_negara_impor(prs, data)
        slide_komoditas_impor(prs, data)
        slide_penutup(prs, data)

        os.makedirs(REPORT_DIR, exist_ok=True)
        if output is None:
            fname = f"Laporan_EKSIM_{BULAN_ID[b1]}-{BULAN_ID[b2]}_{tahun}_{datetime.now():%Y%m%d_%H%M%S}.pptx"
            output = os.path.join(REPORT_DIR, fname)
        prs.save(output)
        log("info", "PPT tersimpan", {"file": output})

        cur.execute("UPDATE automation_jobs SET status='SUCCESS', finished_at=NOW() WHERE id=%s", (job_id,))
        conn.commit()
        conn.close()
        return output, job_id
    except Exception as e:  # noqa: BLE001
        log("error", str(e))
        cur.execute("UPDATE automation_jobs SET status='FAILED', finished_at=NOW(), error_message=%s WHERE id=%s",
                    (str(e), job_id))
        conn.commit()
        conn.close()
        raise


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--tahun", type=int, default=2026)
    ap.add_argument("--bulan-awal", type=int, default=1)
    ap.add_argument("--bulan-akhir", type=int, default=6)
    ap.add_argument("--output", default=None)
    args = ap.parse_args()
    out, job_id = generate(args.tahun, args.bulan_awal, args.bulan_akhir, args.output)
    print(f"SUKSES: {out} (job {job_id})")


if __name__ == "__main__":
    main()