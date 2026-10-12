#!/usr/bin/env python3
"""PHASE 7 — Test generator + API laporan PPT."""
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from backend.api.app import create_app
from backend.db import get_connection

PASS = 0
FAIL = 0


def check(name, cond, detail=""):
    global PASS, FAIL
    if cond:
        PASS += 1
        print(f"  PASS  {name}")
    else:
        FAIL += 1
        print(f"  FAIL  {name}  {detail}")


def main():
    app = create_app()
    app.config["TESTING"] = True
    c = app.test_client()
    tok = c.post("/api/auth/login", json={"email": "admin@kkp.local", "password": "admin123"}).get_json()["token"]
    H = {"Authorization": f"Bearer {tok}"}

    print("== proteksi ==")
    check("list tanpa token 401", c.get("/api/report/list").status_code == 401)
    check("generate tanpa token 401", c.post("/api/report/generate_ppt", json={}).status_code == 401)

    print("== generate ==")
    r = c.post("/api/report/generate_ppt", json={"tahun": 2026, "bulan_awal": 1, "bulan_akhir": 6}, headers=H)
    check("generate 200", r.status_code == 200, r.get_data(as_text=True)[:200])
    d = r.get_json()
    check("file ada di disk", os.path.isfile(d.get("file", "")), str(d))
    if os.path.isfile(d.get("file", "")):
        from pptx import Presentation
        prs = Presentation(d["file"])
        check("PPT 11 slide", len(prs.slides) == 11, str(len(prs.slides)))
        jenis = sum(1 for s in prs.slides for sh in s.shapes if sh.has_chart)
        check("PPT punya chart", jenis >= 3, str(jenis))
        check("PPT tanpa tabel (template flat)", all(not sh.has_table for s in prs.slides for sh in s.shapes))
        import re as _re

        def _texts(slide):
            out = []
            def _w(shapes):
                for sh in shapes:
                    if sh.has_text_frame:
                        out.append("\n".join(p.text for p in sh.text_frame.paragraphs))
                    if sh.shape_type == 6:
                        _w(sh.shapes)
            _w(slide.shapes)
            return out

        s1 = _texts(prs.slides[1])
        s8 = _texts(prs.slides[8])
        check("slide1 langkah nilai+YoY", any(_re.match(r"^USD \d+,\d+ Miliar$", t) for t in s1),
              [t for t in s1 if "Miliar" in t])
        check("slide1 YoY terisi", any(_re.match(r"^\d+,\d+% \(YoY\)$", t) for t in s1))
        ch1 = next(sh.chart for sh in prs.slides[1].shapes if sh.has_chart)
        check("slide1 chart 5 kategori s/d 2026",
              list(ch1.plots[0].categories)[-1:] == ["2026"] and len(ch1.plots[0].categories) == 5,
              str(list(ch1.plots[0].categories)))
        check("slide1 top-5 tanpa catch-all",
              not any(t.split("\n")[0] in ("Lainnya", "Ikan Lainnya") for t in s1),
              [t for t in s1 if t.split("\n")[0] in ("Lainnya", "Ikan Lainnya")])
        check("slide8 TOTAL IMPOR terisi", any("TOTAL IMPOR" in t and "JUTA" in t for t in s8),
              [t for t in s8 if "TOTAL IMPOR" in t])

        s2, s3, s4, s9 = (_texts(prs.slides[i]) for i in (2, 3, 4, 9))
        check("narasi neraca data-driven",
              any("Produk perikanan Indonesia mencapai" in t and "3,42" in t for t in s2),
              [t for t in s2 if "mencapai" in t])
        check("narasi negara tujuan terisi",
              any("masih menjadi negara utama tujuan ekspor" in t for t in s3))
        check("narasi tren komoditas terisi",
              any("Apabila dibandingkan dengan periode yang sama tahun" in t for t in s4))
        joined = "\n".join(s2 + s3 + s4 + s8 + s9)
        check("tidak ada angka narasi template basi",
              not any(x in joined for x in ["1.010,64", "583,71", "60,12 Juta",
                                            "2,62 Miliar", "0,38 Miliar", "Norway"]),
              [x for x in ["1.010,64", "583,71", "60,12 Juta", "2,62 Miliar",
                           "0,38 Miliar", "Norway"] if x in joined])
        allt = [t for i in range(11) for t in _texts(prs.slides[i])]
        check("tidak ada teks artefak cv/v",
              not any(t.strip().lower() in ("cv", "v") for t in allt),
              [t for t in allt if t.strip().lower() in ("cv", "v")])

        import glob as _glob
        import zipfile as _zip
        from lxml import etree as _etree

        _A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
        _bad = 0
        with _zip.ZipFile(d["file"]) as _z:
            for _n in _z.namelist():
                if not (_n.startswith("ppt/slides/slide") and _n.endswith(".xml")):
                    continue
                _root = _etree.fromstring(_z.read(_n))
                for _p in _root.iter(_A + "p"):
                    _seen_end = False
                    for _c in _p:
                        _tag = _c.tag.split("}")[-1]
                        if _tag in ("r", "br", "fld"):
                            if _seen_end:
                                _bad += 1
                        elif _tag == "endParaRPr":
                            _seen_end = True
        check("urutan XML paragraf valid (endParaRPr di akhir)", _bad == 0,
              "pelanggaran urutan: %d" % _bad)

    print("== validasi job tercatat ==")
    conn = get_connection()
    cur = conn.cursor()
    cur.execute("SELECT status FROM automation_jobs WHERE job_type='report_ppt' ORDER BY id DESC LIMIT 1")
    row = cur.fetchone()
    conn.close()
    check("job SUCCESS", row and row["status"] == "SUCCESS", str(row))

    print("== input tidak valid ==")
    check("bulan salah 400", c.post("/api/report/generate_ppt", json={"tahun": 2026, "bulan_awal": 6, "bulan_akhir": 1}, headers=H).status_code == 400)
    check("tahun ngawur 400", c.post("/api/report/generate_ppt", json={"tahun": 99, "bulan_awal": 1, "bulan_akhir": 6}, headers=H).status_code == 400)

    print("== list ==")
    r = c.get("/api/report/list", headers=H)
    check("list 200", r.status_code == 200)
    d = r.get_json()
    check("list ada isi", len(d["reports"]) >= 1)
    check("report punya file valid", any(x["downloadable"] for x in d["reports"]))

    print("== download ==")
    jid = d["reports"][0]["id"]
    r = c.get(f"/api/report/download?job_id={jid}", headers=H)
    check("download 200", r.status_code == 200, str(r.status_code))
    check("content-type pptx", "presentation" in r.headers.get("Content-Type", "") or "octet-stream" in r.headers.get("Content-Type", ""), r.headers.get("Content-Type", ""))
    check("download job ngawur 404", c.get("/api/report/download?job_id=999999", headers=H).status_code == 404)

    print()
    print(f"RESULT: {PASS} passed, {FAIL} failed")
    return 1 if FAIL else 0


if __name__ == "__main__":
    sys.exit(main())